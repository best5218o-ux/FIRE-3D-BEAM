# -*- coding: utf-8 -*-
"""표준양식 정보카드 생성 — templates/소방활동정보카드_표준양식.docx 를 실제로 채워(python-docx) LibreOffice 로 PDF 변환.

양식 항목명(①~⑪)을 그대로 키로 쓰는 정답 `gt['card_std']` = [동별 {관리번호, 주용도, 상호, 소재지{지번,도로명}, 거리{소방서,관할센터,최단센터},
소방용수[{종류,관서,동,번호,거리km}], 건물구조{양식,구조,지붕,동명,지하층수,지상층수,연면적,바닥면적,특별피난계단수,피난계단수,일반계단수,옥외계단수,
승강기_비상,승강기_일반}, 관계인[{구분,성명,연락처}], 위험물[{구분,시설명,품명,수량}], 층별현황[{층,층_원문,용도,면적,소방시설_원문,기타}]}]
+ `gt['card_std_figs']` = [{동, 종류(위치도|사진|진압작전도|평면도), 캡션, floors}]

모든 값은 가상(가상시 가상구, 가상소유자, 000-xxxx-xxxx). 관계인 성명·연락처는 개인정보 항목 → 엔진 기본값은 출력에서 제외(fire3dbeam/card.py redact).
variant: 'std'(LibreOffice 기본 글꼴) | 'hwp'(한글→PDF 흉내: 명조 글꼴·줄간격 160%·자간·'지상1층' 표기·㎡ 생략)"""
import copy, io, os, random, re, shutil, subprocess, tempfile
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, 'templates', '소방활동정보카드_표준양식.docx')
ROLES = ['소유자', '관리자', '소방안전관리자', '위험물안전관리자']
HAZ_ROWS = ['위험물시설', '위험물시설', '가스취급현황', '특수가연물', '기타']


def floor_txt(fid, hwp=False):
    """층 id → 양식 표기(「지하1」「지상1」「옥탑1」). hwp 변형은 '지하1층'처럼 층을 붙인다."""
    if fid == 'RF': t = '옥탑1'
    elif fid.startswith('B'): t = f'지하{fid[1:]}'
    else: t = f'지상{fid[:-1]}'
    return t + ('층' if hwp else '')


def _area(f): return sum((r[2] - r[0]) * (r[3] - r[1]) for r in f['rects'])


# ───────────────────────────── 정답 값(양식 항목명 그대로)
def fields(b, card, rng, seed, dong='주동', multi=False, fp_emev=None):
    """건물 b + 기존 card 값(면적·소방시설·계단 — 기존 정답과 같은 값) → 표준양식 ①~⑩ 정답 dict"""
    fl = sorted(b.floors.values(), key=lambda f: f['level'])
    nb = sum(1 for f in fl if f['level'] < 0); nu = sum(1 for f in fl if f['level'] >= 0)
    no = int(re.sub(r'\D', '', b.name)[-4:] or seed) % 10000
    st = card['stairs']
    gfa = round(sum(card['areas'].values()), 2)
    hall = b.btype == 'hall'
    struct = rng.choice(['일반철골구조', '철골조', '철근콘크리트조']) if hall else rng.choice(['철근콘크리트조', '철근콘크리트조', '철골철근콘크리트조'])
    roof = rng.choice(['샌드위치패널', '칼라강판']) if hall and '콘크리트' not in struct else rng.choice(['슬라브', '평슬라브'])
    d_c = round(rng.uniform(0.6, 4.5), 1)
    d = dict(관리번호=f'가상-{2026}-{no:05d}' + (f'-{dong}' if multi else ''), 주용도=b.use, 상호=b.name,
             소재지=dict(지번=f'가상시 가상구 가상동 {no % 97 + 1}-{no}', 도로명=f'가상시 가상구 가상로{no % 40 + 1}길 {no % 200 + 1}'),
             거리=dict(소방서=round(d_c + rng.uniform(0.5, 4.0), 1), 관할센터=d_c, 최단센터=round(d_c * rng.uniform(0.6, 1.0), 1)),
             소방용수=[dict(종류='지상식', 관서='가상', 동='가상동', 번호='1', 거리km=0.15), dict(종류='지하식', 관서='가상', 동='가상동', 번호='2', 거리km=0.32)],
             건물구조=dict(양식='양식', 구조=struct, 지붕=roof, 동명=dong, 지하층수=nb, 지상층수=nu, 연면적=gfa,
                       바닥면적=round(max(v for k, v in card['areas'].items() if k != 'RF'), 2),
                       특별피난계단수=st['special'], 피난계단수=st['evac'], 일반계단수=st['normal'], 옥외계단수=st['outdoor'],
                       승강기_비상=st['ev_em'], 승강기_일반=st['ev']))
    if rng.random() < 0.3:              # 소방용수 셋째(부지 자료에 없는 번호 → 거리 비교 없음)
        d['소방용수'].append(dict(종류=rng.choice(['급수탑', '비상소화장치']), 관서='가상', 동='가상동', 번호=str(rng.randint(10, 99)), 거리km=round(rng.uniform(0.2, 0.6), 3)))
    # ⑨ 위험물 시설 등
    haz = []
    if card.get('hazmat'):
        haz.append(dict(구분='위험물시설', 시설명='옥내탱크저장소(제4류)', 품명='제2석유류(경유)', 수량='1,000ℓ'))
    if rng.random() < 0.25:
        haz.append(dict(구분='위험물시설', 시설명=rng.choice(['일반취급소(제4류)', '옥외탱크저장소(제4류)']), 품명=rng.choice(['제4석유류(윤활유)', '제1석유류(휘발유)']),
                        수량=f'{rng.choice([2, 4, 10])},000ℓ'))
    if rng.random() < 0.35:
        g = rng.choice([('LPG 저장(용기)', '액화석유가스', '150kg'), ('도시가스 정압기', '도시가스(LNG)', '-'), ('의료가스 저장', '산소', '1,200㎥')])
        haz.append(dict(구분='가스취급현황', 시설명=g[0], 품명=g[1], 수량=g[2]))
    if hall and rng.random() < 0.5:
        haz.append(dict(구분='특수가연물', 시설명='창고 적재', 품명=rng.choice(['합성수지류(발포)', '면화류', '고무류·플라스틱류']), 수량=f'{rng.randint(2, 30)},000kg'))
    if rng.random() < 0.1:
        haz.append(dict(구분='기타', 시설명='비상발전기 연료탱크', 품명='경유', 수량='500ℓ'))
    d['위험물'] = haz
    rel = [dict(구분='소유자', 성명=rng.choice(['가상소유자', '(주)가상개발', '가상홍길동']), 연락처=f'000-{rng.randint(0, 9999):04d}-{rng.randint(0, 9999):04d}'),
           dict(구분='관리자', 성명='가상관리자', 연락처=f'000-{rng.randint(0, 9999):04d}-{rng.randint(0, 9999):04d}'),
           dict(구분='소방안전관리자', 성명=rng.choice(['가상안전관리자', '가상김안전']), 연락처=f'000-{rng.randint(0, 9999):04d}-{rng.randint(0, 9999):04d}')]
    if any(h['구분'] == '위험물시설' for h in haz):
        rel.append(dict(구분='위험물안전관리자', 성명='가상위험물관리자', 연락처=f'000-{rng.randint(0, 9999):04d}-{rng.randint(0, 9999):04d}'))
    d['관계인'] = rel
    # ⑩ 층별 현황(기존 card 면적·소방시설 그대로)
    rows = []
    ids = sorted(card['areas'], key=lambda f: 999 if f == 'RF' else b.floors[f]['level'])
    order = list(reversed(ids)) if rng.random() < 0.5 else ids
    for fid in order:
        use = b.floors[fid]['use'] if fid in b.floors else '계단실·기계실'
        fac = card['fac'][fid]
        rows.append(dict(층=fid, 용도=use, 면적=card['areas'][fid], 소방시설_원문=' '.join(f'{k}:{v}' for k, v in fac.items()),
                         기타=rng.choice(['', '', '', '무창층' if fid.startswith('B') else '', '피난층' if fid == '1F' else ''])))
    d['층별현황'] = rows
    return d


def rf_row(b, card, rng):
    """옥탑 행(정보카드에만 있는 층 — 도면 없음) 추가: card['areas']['RF'], card['fac']['RF']"""
    top = max(b.floors.values(), key=lambda f: f['level'])
    card['areas']['RF'] = round(_area(top) * rng.uniform(0.08, 0.15), 2)
    card['fac']['RF'] = {'유도등': rng.randint(1, 4), '수동식소화기': rng.randint(1, 3), '비상조명등': rng.randint(1, 4)}


# ───────────────────────────── 그림(위치도·사진·진압작전도 자리표시)
def img_location(d, rng):
    from synth.render import font
    W, H = 1000, 620
    im = Image.new('RGB', (W, H), (236, 234, 226)); dr = ImageDraw.Draw(im)
    for x in range(40, W, 140): dr.rectangle([x, 0, x + 18, H], fill=(255, 255, 255))
    for y in range(30, H, 120): dr.rectangle([0, y, W, y + 18], fill=(255, 255, 255))
    cx, cy = 180 + rng.randint(-40, 40), 520; bx, by = 820 + rng.randint(-40, 40), 110
    dr.polygon([(bx - 30, by - 22), (bx + 30, by - 22), (bx + 30, by + 22), (bx - 30, by + 22)], fill=(220, 60, 60))
    dr.ellipse([cx - 16, cy - 16, cx + 16, cy + 16], fill=(30, 90, 200))
    pts = [(cx, cy), (cx, 390), (600, 390), (600, 150), (bx, 150), (bx, by + 22)]
    dr.line(pts, fill=(230, 120, 0), width=7)
    dr.text((cx + 22, cy - 12), '가상119안전센터(관할)', fill=(0, 0, 0), font=font(22))
    dr.text((bx - 110, by - 60), d['상호'], fill=(0, 0, 0), font=font(22))
    dr.text((30, 20), f"관할 119안전센터 → 대상물 {d['거리']['관할센터']:.1f} km · 약 {max(2, round(d['거리']['관할센터'] * 2.4))}분", fill=(0, 0, 0), font=font(24))
    return im


def img_photo(b, rng, side=False):
    from synth.render import font
    W, H = 900, 600
    im = Image.new('RGB', (W, H), (190, 212, 235)); dr = ImageDraw.Draw(im)
    dr.rectangle([0, 470, W, H], fill=(120, 120, 118))
    nu = max(1, sum(1 for f in b.floors.values() if f['level'] >= 0))
    fh = min(40, 400 // nu); w = 360 if side else 560; x0 = (W - w) // 2; y1 = 470; y0 = y1 - fh * nu
    dr.rectangle([x0, y0, x0 + w, y1], fill=(205, 200, 190), outline=(90, 90, 90), width=3)
    for k in range(nu):
        for x in range(x0 + 20, x0 + w - 30, 50):
            dr.rectangle([x, y0 + k * fh + 8, x + 30, y0 + k * fh + fh - 8], fill=(70, 90, 120))
    if not side: dr.rectangle([W // 2 - 30, y1 - 50, W // 2 + 30, y1], fill=(40, 40, 40))
    dr.text((20, 16), '(가상 건물 사진 — 합성)', fill=(40, 40, 40), font=font(20))
    return im


def img_sitemap_placeholder(d):
    from synth.render import font
    im = Image.new('RGB', (900, 600), (228, 228, 228)); dr = ImageDraw.Draw(im)
    dr.rectangle([6, 6, 893, 593], outline=(120, 120, 120), width=4)
    dr.text((250, 270), '진압작전도 (가상 자리표시)', fill=(90, 90, 90), font=font(34))
    dr.text((230, 320), 'synth.make.sitemap_image(b) 훅이 이미지를 주면 대체', fill=(110, 110, 110), font=font(20))
    return im


# ───────────────────────────── docx 채우기
def _body_items(doc):
    return [(el, Paragraph(el, doc) if el.tag == qn('w:p') else Table(el, doc)) for el in doc.element.body.iterchildren() if el.tag in (qn('w:p'), qn('w:tbl'))]


def _font(run, size=None, bold=None, color=(0, 0, 0), ea=None):
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = RGBColor(*color)
    if ea:
        rpr = run._r.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia'): rf.set(qn(k), ea)


def set_cell(cell, text, size=8.5, align=None, bold=False):
    ps = cell.paragraphs
    for p in ps[1:]: p._p.getparent().remove(p._p)
    p = ps[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    lines = str(text).split('\n')
    for i, ln in enumerate(lines):
        r = p.add_run(ln); _font(r, size, bold)
        if i + 1 < len(lines): r.add_break()
    if align: p.alignment = align


def _fmt_num(v, dec=2): return f'{v:,.{dec}f}'


def fill_dong(tabs, d, hwp=False):
    """tabs = ①~⑩ 표 7개(T0..T6) → 값 채움"""
    C = WD_ALIGN_PARAGRAPH.CENTER
    t0, t1, t2, t3, t4, t5, t6 = tabs
    r = t0.rows[0].cells
    set_cell(r[1], d['관리번호'], align=C); set_cell(r[3], d['주용도'], align=C); set_cell(r[5], d['상호'], align=C)
    c = t1.rows[0].cells
    p0, p1 = c[1].paragraphs[:2]
    for p, k in ((p0, '지번'), (p1, '도로명')):
        rr = p.add_run(' ' + d['소재지'][k]); _font(rr, 8.5)
        for x in p.runs[:-1]: x.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    D = d['거리']
    for p, v in zip(c[3].paragraphs[:3], (D['소방서'], D['관할센터'], D['최단센터'])):
        runs = p.runs
        suf = re.sub(r'^\s*Km', 'Km', runs[-1].text) if len(runs) > 1 else 'Km'
        runs[-1].text = f'  {v:.1f} ' + suf
        for x in runs: x.font.color.rgb = RGBColor(0, 0, 0)
    w = '  /  '.join(f"{x['종류']}-{x['관서']}-{x['동']}-{x['번호']} : {x['거리km']:.3f}Km" for x in d['소방용수'])
    set_cell(t2.rows[0].cells[1], w)
    S = d['건물구조']
    vals = [[S['양식'], S['구조'], S['지붕'], S['동명']],
            [S['지하층수'], S['지상층수'], _fmt_num(S['연면적']), _fmt_num(S['바닥면적'])],
            [S['특별피난계단수'], S['피난계단수'], S['일반계단수'], S['옥외계단수']],
            [S['승강기_비상'], S['승강기_일반']]]
    for i, row in enumerate(vals):
        cells = t3.rows[i].cells
        for j, v in enumerate(row): set_cell(cells[2 + 2 * j], v, align=C)
    for rel in d['관계인']:
        i = ROLES.index(rel['구분'])
        cells = t4.rows[i].cells
        set_cell(cells[3], rel['성명']); set_cell(cells[5], rel['연락처'])
    # ⑨: 구분별 행(위험물시설 2행 기본, 모자라면 복제)
    by = {}
    for h in d['위험물']: by.setdefault(h['구분'], []).append(h)
    nh = len(by.get('위험물시설', []))
    for _ in range(max(0, nh - 2)):
        t5.rows[2]._tr.addnext(copy.deepcopy(t5.rows[2]._tr))
    rows = t5.rows[1:]
    k = {'위험물시설': 0}
    for row in rows:
        lab = re.sub(r'\s+', '', row.cells[1].text)
        lst = by.get(lab, [])
        j = k.get(lab, 0)
        if j < len(lst):
            h = lst[j]; k[lab] = j + 1
            set_cell(row.cells[2], h['시설명']); set_cell(row.cells[3], h['품명']); set_cell(row.cells[4], h['수량'], align=C)
    # ⑩: 행 수 맞추기(모자라면 마지막 행 복제, 남으면 삭제)
    fr = d['층별현황']
    while len(t6.rows) - 1 < len(fr): t6.rows[-1]._tr.addnext(copy.deepcopy(t6.rows[-1]._tr))
    for row in list(t6.rows)[1 + len(fr):]: row._tr.getparent().remove(row._tr)
    for row, x in zip(t6.rows[1:], fr):
        cells = row.cells
        set_cell(cells[1], floor_txt(x['층'], hwp), align=C); set_cell(cells[2], x['용도'], align=C)
        set_cell(cells[3], _fmt_num(x['면적']) if hwp else _fmt_num(x['면적']) + '㎡', align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell(cells[4], x['소방시설_원문'], size=8); set_cell(cells[5], x['기타'], align=C)


def _add_picture(tab, im, max_w=172, max_h=118):
    buf = io.BytesIO(); im.save(buf, 'PNG'); buf.seek(0)
    w, h = im.size
    k = min(max_w / w, max_h / h)
    p = tab.rows[0].cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(buf, width=Mm(w * k))


def build_docx(dongs, out_path, variant='std', rng=None):
    """dongs = [{'fields': d, 'figs': [(kind, caption, PIL.Image)]}] → out_path(.docx)
    kind: 위치도|사진|진압작전도|평면도. 여러 동이면 ①~⑩ 과 ⑪ 을 동별로 반복한다(양식 주석 그대로)."""
    rng = rng or random.Random(0)
    doc = Document(TEMPLATE)
    hwp = variant == 'hwp'
    items = _body_items(doc)
    tabs = [x for _, x in items if isinstance(x, Table)]
    paras = [x for _, x in items if isinstance(x, Paragraph)]
    paras[0].runs[0].text = paras[0].runs[0].text.replace('출력일시 :', '출력일시 : 2026-09-29 10:00:00').replace('출력자 :', '출력자 : 가상사용자')
    # 1. 건물개황 블록(①~⑩ 표 + 사이 빈 문단 + 주석 3줄): 첫 표 ~ '※ 동이 여러 개이면' 문단
    body = doc.element.body
    els = list(body.iterchildren())
    i0 = els.index(tabs[0]._tbl)
    i1 = next(i for i, e in enumerate(els) if e.tag == qn('w:p') and '동이 여러 개이면' in Paragraph(e, doc).text)
    block1 = els[i0:i1 + 1]
    # ⑪ 블록: '⑪ 동별' 문단 ~ 마지막 평면도 표(평면도 1개 틀만 남기고 복제)
    j0 = next(i for i, e in enumerate(els) if e.tag == qn('w:p') and Paragraph(e, doc).text.strip().startswith('⑪'))
    jn = next(i for i, e in enumerate(els) if e.tag == qn('w:p') and '평면도 제목은' in Paragraph(e, doc).text)
    fig_els = els[j0:jn]
    # 틀 분해: 위치도(캡션·설명·표), 정면사진(캡션·설명·표), 진압작전도(캡션·설명·표·범례 캡션·범례 표), 평면도 틀
    def find_p(txt, start=0):
        return next(i for i, e in enumerate(fig_els) if i >= start and e.tag == qn('w:p') and txt in Paragraph(e, doc).text)
    ip_loc, ip_ph, ip_sm, ip_p1 = find_p('0. 위치도'), find_p('00 정면사진'), find_p('0 진압작전도'), find_p('1. 평면도')
    ip_p2 = find_p('2. 평면도')
    head = fig_els[:ip_loc]                          # ⑪ 문단
    loc_blk, ph_blk, sm_blk, plan_blk = fig_els[ip_loc:ip_ph], fig_els[ip_ph:ip_sm], fig_els[ip_sm:ip_p1], fig_els[ip_p1:ip_p2]
    anchor = els[jn]                                 # 평면도 주석(끝) 앞에 차례로 넣는다
    for e in fig_els: body.remove(e)
    blocks1 = [block1] + [[copy.deepcopy(e) for e in block1] for _ in dongs[1:]]
    # 동별 ①~⑩: 두 번째 동부터 첫 블록 뒤에 복제
    last = block1[-1]
    for bl in blocks1[1:]:
        sep = OxmlElement('w:p'); last.addnext(sep); Paragraph(sep, doc).add_run().add_break(WD_BREAK.PAGE)     # 동마다 새 쪽
        last = sep
        for e in bl: last.addnext(e); last = e
    for k, (bl, dg) in enumerate(zip(blocks1, dongs)):
        tt = [Table(e, doc) for e in bl if e.tag == qn('w:tbl')]
        fill_dong(tt[:7], dg['fields'], hwp)
    # ⑪ 동별 그림
    def clone(bl): return [copy.deepcopy(e) for e in bl]
    def put(bl, caption=None, im=None, first_p_text=None):
        tb = None
        for e in bl:
            anchor.addprevious(e)
            if e.tag == qn('w:p'): Paragraph(e, doc).paragraph_format.keep_with_next = True      # 캡션·설명이 그림과 다른 쪽으로 떨어지지 않게
            if e.tag == qn('w:tbl') and tb is None: tb = Table(e, doc)
        if caption is not None:
            p = Paragraph(bl[0], doc)
            runs = p.runs
            runs[0].text = caption
            for x in runs[1:]: x.text = ''
        if im is not None and tb is not None: _add_picture(tb, im)
    for k, dg in enumerate(dongs):
        h = clone(head)
        if len(dongs) > 1:
            p = Paragraph(h[0], doc); p.runs[-1].text = p.runs[-1].text + f" — 동명: {dg['fields']['건물구조']['동명']}"
            if k > 0:
                br = OxmlElement('w:p'); anchor.addprevious(br); Paragraph(br, doc).add_run().add_break(WD_BREAK.PAGE)
        for e in h: anchor.addprevious(e)
        n_plan = 0
        for kind, cap, im in dg['figs']:
            if kind == '위치도': put(clone(loc_blk), cap, im)
            elif kind == '사진': put(clone(ph_blk), cap, im)
            elif kind == '진압작전도': put(clone(sm_blk), cap, im)
            else:
                put(clone(plan_blk), cap, im); n_plan += 1
    if hwp: hwpize(doc, rng)                         # 한글→PDF 흉내(한·영 자동 간격 켜 둠 → 글꼴 경계 공백이 끼어듦)
    else:                                            # 기본: 한·영/숫자 자동 간격 끔(LibreOffice 가 공백 글자를 끼워 넣지 않게)
        for p in doc.element.body.iter(qn('w:p')):
            ppr = p.find(qn('w:pPr'))
            if ppr is None: ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
            for k in ('w:autoSpaceDE', 'w:autoSpaceDN'):
                e = OxmlElement(k); e.set(qn('w:val'), '0'); ppr.append(e)
    doc.save(out_path)


def hwpize(doc, rng):
    """한글(HWP)→PDF 출력 흉내: 명조 글꼴, 줄간격 160%, 좁은 자간, 쪽 넘김 표 머리행 반복 없음, 한·영 자동 간격(끼어든 공백)"""
    ea = 'Noto Serif CJK KR'
    for r in doc.element.body.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None: rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
        rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for k in ('w:ascii', 'w:hAnsi', 'w:eastAsia'): rf.set(qn(k), ea)
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(rng.choice([-6, -4, 0])))
        rpr.append(sp)
    for h in list(doc.element.body.iter(qn('w:tblHeader'))): h.getparent().remove(h)      # 쪽 넘김 표에 머리행 반복 없음
    for p in doc.element.body.iter(qn('w:p')):
        ppr = p.find(qn('w:pPr'))
        if ppr is None: ppr = OxmlElement('w:pPr'); p.insert(0, ppr)
        s = ppr.find(qn('w:spacing'))
        if s is None: s = OxmlElement('w:spacing'); ppr.append(s)
        s.set(qn('w:line'), '384'); s.set(qn('w:lineRule'), 'auto')


def to_pdf(pairs, timeout=900):
    """[(docx 경로, pdf 경로)] → LibreOffice headless 한 번에 변환(이미지는 무손실 — 도면 화소 보존)"""
    if not pairs: return
    tmp = tempfile.mkdtemp(prefix='stdcard_')
    try:
        srcs = []
        for i, (dx, _) in enumerate(pairs):
            s = os.path.join(tmp, f'c{i:04d}.docx'); shutil.copy(dx, s); srcs.append(s)
        filt = 'pdf:writer_pdf_Export:{"UseLosslessCompression":{"type":"boolean","value":"true"},"ReduceImageResolution":{"type":"boolean","value":"false"}}'
        env = dict(os.environ, HOME=tmp)
        for k in range(0, len(srcs), 20):
            subprocess.run(['soffice', '--headless', '--norestore', f'-env:UserInstallation=file://{tmp}/lo', '--convert-to', filt, '--outdir', tmp] + srcs[k:k + 20],
                           check=True, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=timeout)
        for i, (_, pdf) in enumerate(pairs):
            shutil.move(os.path.join(tmp, f'c{i:04d}.pdf'), pdf)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
